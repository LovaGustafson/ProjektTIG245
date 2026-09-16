"""Read source Excel data into an independent in-memory DataFrame."""

from itertools import chain, islice
from os import PathLike
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.mapping.column_mapper import COLUMN_MAPPING


class ExcelReadError(ValueError):
    """The supplied file could not be parsed as an Excel workbook."""


HEADER_SCAN_ROWS = 50
HEADER_FIELDS = {'Vernr', 'Vrad', 'Verdatum', 'Utfall', 'Konto', 'Vertyp',
                 'Huvudtext', 'Radtext'}
HEADER_ALIASES = {alias: COLUMN_MAPPING[source]
                  for source in HEADER_FIELDS
                  for alias in (source, COLUMN_MAPPING[source])}


def _header_position(preview):
    # Count distinct fields, not repeated labels or substrings in metadata text.
    scores = [len({HEADER_ALIASES[value.strip()] for value in row
                   if isinstance(value, str) and value.strip() in HEADER_ALIASES})
              for row in preview]
    best = max(scores, default=0)
    # Keep first-row reading for generic reports and small/incomplete schemas.
    return scores.index(best) if best >= 3 else 0


def _prepare_columns(data):
    """Remove only unnamed empty columns; retain every named or populated one."""
    def blank(value):
        return pd.isna(value) or (isinstance(value, str) and not value.strip())

    positions = [i for i, name in enumerate(data.columns)
                 if not blank(name) or not data.iloc[:, i].map(blank).all()]
    result = data.iloc[:, positions].copy(deep=True)
    reserved = {str(name).strip() for name in result.columns if not blank(name)}
    names = []
    for position, name in zip(positions, result.columns):
        if blank(name):
            name = f'Namnlös kolumn {position + 1}'
            while name in reserved:
                name = '_' + name
            reserved.add(name)
        names.append(name)
    result.columns = pd.Index(names, dtype=object)
    return result


def read_excel(
    path: str | PathLike[str], *, sheet_name: str | int = 0
) -> pd.DataFrame:
    """Read every row and column of one .xlsx worksheet without writing to disk.

    The first worksheet is the default; select others by name or index.
    Read up to 50 rows without interpreting headers (equivalent to header=None).
    The earliest row with the most distinct schema fields (at least three) is
    the header. Source and standardized field names are accepted, with outer
    whitespace ignored. If none qualifies, retain first-row reading for generic
    reports and incomplete schemas. Sheets are never combined.

    Metadata above the header is skipped. Unnamed empty columns are removed;
    unnamed populated columns receive unique names. Named duplicate headers
    remain available to validation rather than silently changing their meaning.
    Original named headers, blank cells, text identifiers,
    Excel error values and formula expressions are retained. Excel-native
    dates/numbers are decoded by openpyxl; display formatting is not data.
    Formulas are not evaluated. No business validation or transformation runs.

    Returns a new object-typed DataFrame, independent of the source file.
    Missing files, directories and permissions raise their standard OSError
    subclasses. Unsupported extensions or sheet selectors raise ValueError
    (TypeError for invalid selector types). Unreadable workbooks raise
    ExcelReadError with the underlying exception preserved as the cause.
    """
    source = Path(path)
    if source.suffix.lower() != ".xlsx":
        raise ValueError(f"Expected an .xlsx file: {source}")
    if isinstance(sheet_name, bool) or not isinstance(sheet_name, (str, int)):
        raise TypeError("sheet_name must be a worksheet name or zero-based index")

    # A binary read-only handle prevents this module from writing to the source.
    with source.open("rb") as stream:
        workbook = None
        try:
            workbook = load_workbook(stream, read_only=True, data_only=False)
            names = workbook.sheetnames
            if isinstance(sheet_name, int):
                if sheet_name < 0 or sheet_name >= len(names):
                    raise ValueError(f"Worksheet index {sheet_name} does not exist")
                selected = names[sheet_name]
            else:
                selected = sheet_name
                if selected not in names:
                    raise ValueError(f"Worksheet {selected!r} does not exist")

            rows = workbook[selected].iter_rows(values_only=True)
            preview = list(islice(rows, HEADER_SCAN_ROWS))
            if not preview:
                return pd.DataFrame(dtype=object)
            header_position = _header_position(preview)
            headers = preview[header_position]
            data = pd.DataFrame(
                chain(preview[header_position + 1:], rows),
                columns=pd.Index(headers, dtype=object), dtype=object,
            )
            return _prepare_columns(data)
        except OSError:
            raise
        except Exception as exc:
            # Workbook parsing can fail with ZIP, XML or engine-specific errors.
            # Keep this boundary local to ingestion and retain the original cause.
            raise ExcelReadError(
                f"Could not read .xlsx file {source} (sheet {sheet_name!r}): {exc}"
            ) from exc
        finally:
            if workbook is not None:
                workbook.close()
