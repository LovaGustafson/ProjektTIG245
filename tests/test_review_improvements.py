"""Synthetic source-to-export and dashboard regression coverage."""
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest
from streamlit.proto.WidgetStates_pb2 import WidgetStates

from src.ui_support import analyze_upload, filter_details, review_row_detail
from src.output.report_generator import review_tables

ROOT = Path(__file__).resolve().parents[1]


def source(date_header='VerDat'):
    book = Workbook()
    sheet = book.active
    sheet.append(['Syntetisk rapport'])
    sheet.append(['Vernr', 'Vrad', date_header, 'Konto', 'Vertyp', 'Utfall',
                  'Mm', 'Pg', 'Radtext', 'Leverantör', 'Huvudtext'])
    for i, (account, kind) in enumerate([('7698', 'X'), ('4000', 'X'),
                                        ('7699 ', 'KR01'), ('4000', 'KR01'), ('4000', 'X')]):
        sheet.append(['001', 1, None if i == 4 else datetime(2026, 9, 8, 12, 30),
                      account, kind, 'fel' if i == 1 else 10,
                      [None, '001', 3, 'NA', None][i],
                      ['007', 2, 'Pg text', '#N/A', None][i],
                      ['  Text med blanksteg  ', 'Radtext\nandra raden', '=literal', 'NA', None][i],
                      'Testleverantör', 'Huvudtext'])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


@pytest.mark.parametrize('date_header', ['Verdatum', 'VerDat', 'VerDatum'])
def test_source_fields_survive_all_exports_and_pipeline(tmp_path, date_header):
    path = tmp_path / 'source.xlsx'
    path.write_bytes(source(date_header))
    before, mtime = path.read_bytes(), path.stat().st_mtime_ns
    review = analyze_upload(before)
    result = review.result
    original = result.original_data.copy(deep=True)
    standardized = result.standardized_data.copy(deep=True)
    for source_field, internal in [(date_header, 'verification_date'), ('Mm', 'vat_code'),
                                   ('Pg', 'posting_group'), ('Radtext', 'line_text')]:
        assert original[source_field].tolist() == standardized[internal].tolist()
    kept, excluded = review_tables(original, result.filtering)
    for filename, sheets in [('granskning.xlsx', {'Granskning': kept}),
                             ('bortfiltrerade.xlsx', {'Bortfiltrerade': excluded}),
                             ('samlad_kontrollfil.xlsx', {'Granskning': kept, 'Bortfiltrerade': excluded})]:
        book = load_workbook(BytesIO(review.downloads[filename]))
        for sheet_name, expected in sheets.items():
            sheet = book[sheet_name]
            assert list(next(sheet.values)) == list(expected.columns)
            assert list(sheet.values)[1:] == list(expected.itertuples(index=False, name=None))
            date_column = list(expected.columns).index(date_header) + 1
            for i, value in enumerate(expected[date_header], 2):
                if value is not None:
                    assert sheet.cell(i, date_column).data_type == 'd'
        book.close()
    pd.testing.assert_frame_equal(result.original_data, original)
    pd.testing.assert_frame_equal(result.standardized_data, standardized)
    assert path.read_bytes() == before
    assert path.stat().st_mtime_ns == mtime


def test_filter_transparency_counts_rows_and_overlap():
    result = analyze_upload(source()).result
    counts, rows = filter_details(result, 'account')
    assert counts.to_dict('list') == {'Konto': ['7698', '7699'], 'Antal rader': [1, 1]}
    assert rows.index.tolist() == [0, 2]
    counts, rows = filter_details(result, 'verification_type')
    assert counts.set_index('Vertyp').loc['KR01', 'Antal rader'] == 2
    assert rows.index.tolist() == [2, 3]
    updated = analyze_upload(source(), excluded_verification_types=['X']).result
    counts, rows = filter_details(updated, 'verification_type')
    assert counts.to_dict('list') == {'Vertyp': ['X'], 'Antal rader': [3]}
    assert rows.index.tolist() == [0, 1, 4]


def test_selected_row_uses_source_position_and_existing_warnings():
    result = analyze_upload(source()).result
    fields, warnings = review_row_detail(result, 0)
    values = fields.set_index('Fält')['Källvärde']
    assert values['Radtext'] == 'Radtext\nandra raden'
    assert values['Mm'] == '001'
    assert values['Leverantör'] == 'Testleverantör'
    assert any('Utfall' in warning for warning in warnings)
    fields, warnings = review_row_detail(result, 1)
    assert fields.set_index('Fält').loc['Radtext', 'Källvärde'] is None
    assert any('Verdatum' in warning for warning in warnings)


def test_dashboard_full_labels_transparency_and_selected_detail():
    app = AppTest.from_file(str(ROOT / 'streamlit_app.py'))
    app.session_state['review'] = analyze_upload(source())
    app.run()
    assert not app.exception
    assert [b.label.split('\n')[1] for b in app.button if b.key and b.key.startswith('kpi_')] == [
        'Totalt antal rader', 'Kvar för granskning', 'Exkluderade på grund av konto',
        'Exkluderade på grund av verifikationstyp', 'Totalt bortfiltrerade']
    css = (ROOT / '.streamlit/style.css').read_text()
    assert '.st-key-dashboard_kpis button p' in css
    for rule in ['white-space: normal', 'text-overflow: clip', 'overflow-wrap: anywhere']:
        assert rule in css
    app.button(key='kpi_account').click().run()
    account = app
    # The opened KPI panel precedes the tab tables.
    assert account.dataframe[0].value['Antal rader'].tolist() == [1, 1]
    assert account.dataframe[1].value['Konto'].tolist() == ['7698', '7699 ']
    app.button(key='kpi_verification_type').click().run()
    types = app
    assert types.dataframe[1].value['Vertyp'].tolist() == ['KR01', 'KR01']
    table = app.tabs[0].dataframe[0]
    # AppTest has no public dataframe click helper; send the actual widget event.
    states = WidgetStates()
    state = states.widgets.add()
    state.id = table.proto.id
    state.string_value = json.dumps({'selection': {'rows': [0], 'columns': [], 'cells': []}})
    app._run(states)
    assert not app.exception
    assert any(s.value == 'Vald rad – detaljer' for s in app.subheader)
    assert 'Radtext\nandra raden' in [t.value for t in app.text]
    assert '001' in [t.value for t in app.text]
    assert any('inte träffar någon aktiv filterregel' in i.value and
               'inte automatiskt' in i.value for i in app.info)
    assert any('Utfall' in w.value for w in app.warning)
    # Refiltering invalidates the previous row selection.
    app.multiselect(key='excluded_types').select('X').run()
    assert not app.exception
    assert not any(s.value == 'Vald rad – detaljer' for s in app.subheader)


@pytest.mark.parametrize('date_header', ['VerDat', 'VerDatum'])
def test_date_alias_participates_in_header_detection(tmp_path, date_header):
    from src.ingestion.excel_reader import read_excel
    from src.mapping.column_mapper import map_columns
    book = Workbook()
    sheet = book.active
    sheet.append(['Rapport'])
    sheet.append(['Vernr', 'Vrad', date_header])
    sheet.append(['001', 1, datetime(2026, 9, 8)])
    path = tmp_path / 'header.xlsx'
    book.save(path)
    book.close()
    data = map_columns(read_excel(path))
    assert data['verification_date'].tolist() == [datetime(2026, 9, 8)]


def test_detail_invalid_identity_missing_values_and_no_mutation():
    book = load_workbook(BytesIO(source()))
    book.active['A4'] = 123
    book.active['F4'] = None
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    review = analyze_upload(buffer.getvalue())
    original = review.result.original_data.copy(deep=True)
    downloads = dict(review.downloads)
    fields, warnings = review_row_detail(review.result, 0)
    assert fields.set_index('Fält').loc['Vernr', 'Källvärde'] == 123
    assert any('Vernr' in w and 'format' in w for w in warnings)
    assert any('Utfall' in w and 'saknas' in w for w in warnings)
    pd.testing.assert_frame_equal(review.result.original_data, original)
    assert review.downloads == downloads


@pytest.mark.parametrize('key,title,positions', [
    ('total', 'Totalt antal rader', [0, 1, 2, 3, 4]),
    ('review', 'Kvar för granskning', [1, 4]),
    ('account', 'Exkluderade på grund av konto', [0, 2]),
    ('verification_type', 'Exkluderade på grund av verifikationstyp', [2, 3]),
    ('excluded', 'Totalt bortfiltrerade', [0, 2, 3]),
])
def test_every_kpi_click_opens_correct_source_rows(key, title, positions):
    from src.ui_support import display_dataframe
    review = analyze_upload(source())
    original = review.result.original_data.copy(deep=True)
    downloads = dict(review.downloads)
    app = AppTest.from_file(str(ROOT / 'streamlit_app.py'))
    app.session_state['review'] = review
    app.run()
    assert not app.exception
    assert title not in [heading.value for heading in app.subheader]
    button = app.button(key=f'kpi_{key}')
    assert button.label == f'**{len(positions)}**  \n{title}'
    button.click().run()
    assert not app.exception
    assert app.session_state['selected_kpi'] == key
    assert title in [heading.value for heading in app.subheader]
    row_table = app.dataframe[1 if key in ('account', 'verification_type') else 0].value
    # Streamlit's Arrow round trip infers pandas dtypes; compare displayed values.
    actual = row_table.iloc[:, :len(original.columns)].astype(object)
    expected = display_dataframe(original.iloc[positions]).astype(object)
    pd.testing.assert_frame_equal(actual.where(actual.notna(), None),
                                  expected.where(expected.notna(), None),
                                  check_dtype=False, check_column_type=False)
    if key == 'total':
        assert any('Metadata-rader' in m.value and 'räknas inte' in m.value for m in app.markdown)
    elif key == 'review':
        assert app.dataframe[0].key == 'kpi_review_rows'
        assert any('inte automatiskt' in message.value for message in app.info)
    elif key == 'account':
        assert app.dataframe[0].value.to_dict('list') == {
            'Konto': ['7698', '7699'], 'Antal rader': [1, 1]}
    elif key == 'verification_type':
        from src.filtering.filter_engine import load_exclusions
        counts = app.dataframe[0].value.set_index('Vertyp')['Antal rader']
        assert set(counts.index) == set(load_exclusions()['excluded_verification_types'])
        assert counts['KR01'] == 2
        assert counts.sum() == 2
    else:
        assert row_table['Exkluderingsorsak'].tolist() == [
            'Exkluderad – konto 7698',
            'Exkluderad – konto 7699; verifikationstyp KR01',
            'Exkluderad – verifikationstyp KR01']
    pd.testing.assert_frame_equal(review.result.original_data, original)
    assert review.downloads == downloads


def test_kpi_review_selection_and_filter_changes_use_current_rows():
    app = AppTest.from_file(str(ROOT / 'streamlit_app.py'))
    app.session_state['review'] = analyze_upload(source())
    app.run()
    app.button(key='kpi_review').click().run()
    states = WidgetStates()
    state = states.widgets.add()
    state.id = app.dataframe[0].proto.id
    state.string_value = json.dumps({'selection': {'rows': [0], 'columns': [], 'cells': []}})
    app._run(states)
    assert not app.exception
    assert 'Radtext\nandra raden' in [text.value for text in app.text]
    assert any('Utfall' in warning.value for warning in app.warning)
    app.multiselect(key='excluded_types').select('X').run()
    assert not app.exception
    assert app.dataframe[0].value.empty
    assert not any(s.value == 'Vald rad – detaljer' for s in app.subheader)
    app.button(key='kpi_verification_type').click().run()
    assert app.dataframe[0].value.set_index('Vertyp').loc['X', 'Antal rader'] == 3
    assert len(app.dataframe[1].value) == 5
    # New upload also clears the open detail panel and selection.
    app.file_uploader[0].set_value(('new.xlsx', source(),
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')).run()
    assert 'selected_kpi' not in app.session_state
    assert not any(b.key and b.key.startswith('kpi_') for b in app.button)
