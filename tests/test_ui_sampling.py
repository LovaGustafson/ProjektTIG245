"""Session stability, population changes and exact sample/export linkage."""
from dataclasses import asdict, replace
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from src.sampling.manual_sample import plan_random_sample
from src.ui_support import analyze_upload, select_sample_method, DeferredDownloads
from tests.test_feature_package import invoices, workbook
from tests.synthetic_org import org

APP = Path(__file__).resolve().parents[1] / 'streamlit_app.py'


def population(count):
    data = invoices(count)
    data['Huvudtext'] = [f'Syntetisk leverantör {i} AB Prelb 1' for i in range(count)]
    return data


def sample_ids(review):
    return [v.verification_id for v in review.result.manual_sample]


def start_app(data):
    review = analyze_upload(workbook(data), registry_mode='disabled', defer_downloads=True)
    app = AppTest.from_file(str(APP), default_timeout=30)
    app.session_state['review'] = review
    app.run()
    assert not app.exception
    return app, review


def assert_visible_export(app):
    review = app.session_state['review']
    plan = review.result.sampling_result
    manual = next(tab for tab in app.tabs if tab.label == 'Manuellt stickprov')
    selected = next(table.value for table in manual.dataframe
                    if list(table.value.columns) == ['verification_id', 'supplier', 'supplier_classification'])
    assert selected.verification_id.tolist() == sample_ids(review)
    sheets = pd.read_excel(BytesIO(review.downloads['manual_sample.xlsx']), sheet_name=None, dtype=object)
    assert sheets['rows'].verification_id.tolist() == [
        v.verification_id for v in review.result.manual_sample for _ in v.rows.index]
    assert sheets['Källspårning'].source_row_position.tolist() == [
        p for v in review.result.manual_sample for p in v.rows.index]
    metadata = sheets['Urvalsmetod'].iloc[0]
    assert metadata['Önskat antal verifikationer'] == plan.target_size
    assert metadata['Antal valda verifikationer'] == len(plan.sample)
    assert metadata['Antal verifikationer i populationen'] == len(review.result.verifications)
    assert metadata['Genereringstidpunkt (UTC)'] == plan.generated_at
    assert metadata['Populationens identitetsunderlag SHA-256'] == plan.population_fingerprint
    if plan.method == 'randomized':
        assert metadata['Urvalsmetod'] == 'Slumpmässig'
        assert metadata['Seed'] == str(plan.seed)
        assert pd.isna(metadata['Intervall'])
        assert 'hela' in metadata['Metod'].lower()
        assert sheets['Urvalsbeslut'].draw_position.tolist() == plan.decisions.draw_position.tolist()
    else:
        assert metadata['Urvalsmetod'] == 'Var 20:e (stegvis)'
        assert pd.isna(metadata['Seed'])
        assert metadata['Intervall'] == 20
    counts = sheets['Körningsöversikt'].set_index('Mått').Antal
    assert counts['Verifikationer för manuell granskning'] == len(plan.sample)


def test_method_reruns_view_filters_regeneration_and_exports(monkeypatch):
    seeds = iter([123, 456])
    monkeypatch.setattr('src.ui_support.new_sample_seed', lambda previous=None: next(seeds))
    app, original = start_app(population(85))
    assert app.radio(key='manual_sample_method').value == 'systematic'
    assert app.radio(key='manual_sample_method').options == ['Var 20:e (stegvis)', 'Slumpmässig']
    assert sample_ids(original) == ['20', '40', '60', '80']
    assert_visible_export(app)
    # Method changes and ordinary interactions must not reread/reanalyze sources.
    with patch('src.ui_support.run_pipeline', side_effect=AssertionError('unexpected reanalysis')):
        app.radio(key='manual_sample_method').set_value('randomized').run()
        assert not app.exception
        first = app.session_state['review']
        assert isinstance(first.downloads, DeferredDownloads)
        assert first.result.sampling_result.seed == 123
        assert len(sample_ids(first)) == 20
        assert set(sample_ids(first)) - set(sample_ids(original))
        first_export = first.downloads['manual_sample.xlsx']
        assert_visible_export(app)
        before, after = asdict(original.result.summary), asdict(first.result.summary)
        for field in ('sample_interval', 'sampled_verifications', 'sampled_rows'):
            before.pop(field)
            after.pop(field)
        assert before == after
        for field in ('original_data', 'standardized_data', 'filtering', 'verifications',
                      'detection_results', 'supplier_analysis', 'supplier_view', 'validation'):
            assert getattr(first.result, field) is getattr(original.result, field)

        app.run()
        app.button(key='kpi_review').click().run()
        app.multiselect(key='vf:manual:columns').set_value([0]).run()
        app.text_input(key='vf:manual:c0:text').set_value(sample_ids(first)[0]).run()
        assert not app.exception
        assert app.session_state['review'] is first
        assert_visible_export(app)
        app.button(key='home').click().run()
        assert app.session_state['review'] is first

        app.radio(key='manual_sample_method').set_value('systematic').run()
        assert not app.exception
        assert sample_ids(app.session_state['review']) == sample_ids(original)
        assert_visible_export(app)
        app.radio(key='manual_sample_method').set_value('randomized').run()
        assert not app.exception
        restored = app.session_state['review']
        assert restored.result.sampling_result is first.result.sampling_result
        assert_visible_export(app)

        app.button(key='regenerate_manual_sample').click().run()
        assert not app.exception
        regenerated = app.session_state['review']
        assert regenerated.result.sampling_result.seed == 456
        assert regenerated.result.sampling_result.generated_at != first.result.sampling_result.generated_at
        assert sample_ids(regenerated) != sample_ids(first)
        assert regenerated.downloads is not first.downloads
        assert regenerated.result.sampling_result.population_fingerprint == first.result.sampling_result.population_fingerprint
        assert first.downloads['manual_sample.xlsx'] == first_export
        assert_visible_export(app)
        overview = next(c for c in app.get('flex_container') if c.key == 'sample_overview')
        assert overview.metric[1].value == 'Slumpmässigt'
        assert overview.metric[2].value == '20'
        assert any(m.label == 'Önskat antal fakturor' and m.value == '20' for m in app.metric)


def test_lr_filter_changes_invalidate_random_cache_and_keep_selected_method(monkeypatch):
    seeds = iter([123, 456, 789])
    monkeypatch.setattr('src.ui_support.new_sample_seed', lambda previous=None: next(seeds))
    data = population(44)
    data.loc[:3, 'Vertyp'] = ['LR91S', 'LR95S', 'LR91S', 'FBFM']
    data.loc[2, 'Konto'] = '7698'
    app, initial = start_app(data)
    app.radio(key='manual_sample_method').set_value('randomized').run()
    first = app.session_state['review']
    assert not app.exception
    assert first.result.summary.eligible_verifications == 40
    assert not set(sample_ids(first)) & {'1', '2', '3', '4'}
    pd.testing.assert_frame_equal(first.result.sampling_result.identity_rows, initial.result.sampling_result.identity_rows)
    assert_visible_export(app)
    for value, count, seed in [('LR91S', 41, 456), ('LR95S', 42, 789)]:
        app.multiselect(key='excluded_types').unselect(value).run()
        assert not app.exception
        changed = app.session_state['review']
        assert app.radio(key='manual_sample_method').value == 'randomized'
        assert changed.result.sampling_result.seed == seed
        assert changed.result.summary.eligible_verifications == count
        assert len(changed.result.sampling_result.decisions) == count
        assert changed.downloads is not first.downloads
        assert changed.result.sampling_result.population_fingerprint != first.result.sampling_result.population_fingerprint
        assert changed.source_content == initial.source_content
        assert app.button(key='kpi_review').label.startswith(f'**{count}**')
        assert_visible_export(app)
        assert any('Exkluderingarna ändrades' in i.value for i in app.info)
    app.radio(key='manual_sample_method').set_value('systematic').run()
    assert sample_ids(app.session_state['review']) == ['22', '42']
    assert_visible_export(app)


def test_new_analysis_resets_to_systematic_and_does_not_reuse_old_random():
    app, _ = start_app(population(60))
    app.radio(key='manual_sample_method').set_value('randomized').run()
    previous = app.session_state['review']
    app.file_uploader[0].set_value(('new.xlsx', workbook(population(21)),
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')).run()
    assert 'review' not in app.session_state
    app.button[0].click().run()
    assert not app.exception
    assert app.radio(key='manual_sample_method').value == 'systematic'
    assert sample_ids(app.session_state['review']) == ['20']
    app.radio(key='manual_sample_method').set_value('randomized').run()
    current = app.session_state['review']
    assert not app.exception
    assert len(sample_ids(current)) == 20
    assert current.result.sampling_result.seed != previous.result.sampling_result.seed
    assert len(current.result.sampling_result.decisions) == 21
    assert_visible_export(app)


@pytest.mark.parametrize('count', [0, 1, 7])
def test_ui_empty_or_small_random_population_has_target_actual_and_reason(count):
    app, _ = start_app(population(count))
    app.radio(key='manual_sample_method').set_value('randomized').run()
    assert not app.exception
    assert any(m.label == 'Önskat antal fakturor' and m.value == '20' for m in app.metric)
    assert any(m.label == 'Faktiskt antal fakturor' and m.value == str(count) for m in app.metric)
    assert any(f'{count} av önskade 20' in warning.value for warning in app.warning)
    assert_visible_export(app)


def test_seed_precision_complete_group_provenance_and_all_export_metadata(monkeypatch):
    seed = 2**127 + 123456789
    monkeypatch.setattr('src.ui_support.new_sample_seed', lambda previous=None: seed)
    data = population(4)
    data.loc[1, ['Vernr', 'Huvudtext']] = data.loc[0, ['Vernr', 'Huvudtext']]
    review = select_sample_method(analyze_upload(workbook(data), registry_mode='disabled', defer_downloads=True), 'randomized')
    assert sample_ids(review) == ['1', '3', '4']
    for filename in review.downloads:
        book = load_workbook(BytesIO(review.downloads[filename]))
        metadata = book['Urvalsmetod']
        seed_column = [c.value for c in metadata[1]].index('Seed') + 1
        assert metadata.cell(2, seed_column).value == str(seed)
        assert metadata.cell(2, seed_column).data_type == 's'
        book.close()
    sheets = pd.read_excel(BytesIO(review.downloads['manual_sample.xlsx']), sheet_name=None)
    assert sheets['rows'].verification_id.tolist() == [1, 1, 3, 4]
    assert sheets['Källspårning'].source_row_position.tolist() == [0, 1, 2, 3]
    assert review.result.summary.sampled_verifications == 3
    assert review.result.summary.sampled_rows == 4
    recreated = plan_random_sample(review.result.verifications, seed=seed,
                                   supplier_analysis=review.result.supplier_analysis)
    pd.testing.assert_frame_equal(recreated.decisions, review.result.sampling_result.decisions)


def test_random_uses_existing_strong_aliases_and_is_independent_of_other_results(monkeypatch):
    from src.models.result import CheckResult, CheckStatus
    monkeypatch.setattr('src.ui_support.new_sample_seed', lambda previous=None: 123)
    data = population(6)
    data.Huvudtext = ['Alfa AB Prelb 1', 'Alfa Namn AB Prelb 1', 'Beta AB Prelb 1',
                     ' BETA Aktiebolag. Prelb 1', 'Gamma AB Prelb 1', None]
    register = pd.DataFrame({'Leverantör': ['Alfa AB', 'Alfa Namn AB', 'Beta AB'],
                             'Organisationsnummer': [org(1), org(1), org(2)]})
    review = analyze_upload(workbook(data), registry_mode='uploaded',
                            registry_content=workbook(register), defer_downloads=True)
    first = select_sample_method(review, 'randomized')
    assert len(sample_ids(first)) == 3
    selected = first.result.sampling_result.decisions.query("decision == 'SELECTED'")
    assert set(selected.supplier_key) == {f'org:{org(1)}', f'org:{org(2)}', 'name:gamma ab'}
    # Neither flags nor view classifications enter the selection interface.
    changed = replace(review, result=replace(review.result,
        detection_results=[replace(d, checks=(CheckResult(d.verification_id, 'synthetic',
            CheckStatus.FLAGGED, 'Synthetic finding'),)) for d in review.result.detection_results],
        supplier_view=review.result.supplier_view.assign(excluded_from_view=True)))
    second = select_sample_method(changed, 'randomized')
    assert sample_ids(first) == sample_ids(second)
    pd.testing.assert_frame_equal(first.result.sampling_result.decisions, second.result.sampling_result.decisions)
    pd.testing.assert_frame_equal(first.result.supplier_analysis.rows, review.result.supplier_analysis.rows)
