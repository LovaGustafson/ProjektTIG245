"""Excel Reader tests using only synthetic workbooks in temporary directories."""

from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import Workbook

from src.ingestion.excel_reader import ExcelReadError, read_excel


@pytest.fixture
def invoice_file(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoices"
    sheet.append(["Vernr", "Vrad", "Verdatum", "Utfall", "Konto", "Extra"])
    sheet.append(["00123", 1, datetime(2026, 9, 3), 125.5, "7698", "NA"])
    sheet.append(["00123", 1, "invalid date", "invalid amount", "7699", "NULL"])
    sheet.append([None] * 6)
    sheet.append([None, 2, None, -10, None, "  unchanged  "])
    other = workbook.create_sheet("Other")
    other.append(["Note"])
    other.append(["Synthetic second sheet"])
    workbook.save(path)
    workbook.close()
    return path


def test_reads_all_data_without_mapping_filtering_or_validation(invoice_file):
    result = read_excel(invoice_file)
    expected = pd.DataFrame(
        [
            ["00123", 1, datetime(2026, 9, 3), 125.5, "7698", "NA"],
            ["00123", 1, "invalid date", "invalid amount", "7699", "NULL"],
            [None] * 6,
            [None, 2, None, -10, None, "  unchanged  "],
        ],
        columns=pd.Index(
            ["Vernr", "Vrad", "Verdatum", "Utfall", "Konto", "Extra"], dtype=object
        ),
        dtype=object,
    )
    pd.testing.assert_frame_equal(result, expected)


def test_source_bytes_and_modification_time_unchanged_after_mutation(invoice_file):
    original_bytes = invoice_file.read_bytes()
    original_mtime = invoice_file.stat().st_mtime_ns
    result = read_excel(str(invoice_file))
    result.iloc[0, 0] = "changed in memory"
    result.drop(columns=["Extra"], inplace=True)
    assert invoice_file.read_bytes() == original_bytes
    assert invoice_file.stat().st_mtime_ns == original_mtime
    reread = read_excel(invoice_file)
    assert reread.iloc[0, 0] == "00123"
    assert "Extra" in reread.columns


def test_reads_file_with_read_only_permissions(invoice_file):
    invoice_file.chmod(0o444)
    try:
        assert len(read_excel(invoice_file)) == 4
    finally:
        invoice_file.chmod(0o644)


@pytest.mark.parametrize("selector", ["Other", 1])
def test_selects_worksheet(invoice_file, selector):
    result = read_excel(invoice_file, sheet_name=selector)
    assert result.columns.tolist() == ["Note"]
    assert result.iloc[0, 0] == "Synthetic second sheet"


@pytest.mark.parametrize("selector", ["Missing", 2, -1])
def test_missing_worksheet_has_clear_error(invoice_file, selector):
    with pytest.raises(ValueError, match="Worksheet.*does not exist"):
        read_excel(invoice_file, sheet_name=selector)


@pytest.mark.parametrize("selector", [None, True, 1.5, []])
def test_invalid_worksheet_selector_type(invoice_file, selector):
    with pytest.raises(TypeError, match="sheet_name"):
        read_excel(invoice_file, sheet_name=selector)


def test_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError, match="missing.xlsx"):
        read_excel(tmp_path / "missing.xlsx")


def test_directory_is_not_a_file(tmp_path):
    directory = tmp_path / "directory.xlsx"
    directory.mkdir()
    with pytest.raises((IsADirectoryError, PermissionError)):
        read_excel(directory)


@pytest.mark.parametrize("suffix", [".xls", ".csv", ".txt"])
def test_unsupported_extension(tmp_path, suffix):
    path = tmp_path / f"synthetic{suffix}"
    path.write_text("synthetic data")
    with pytest.raises(ValueError, match="Expected an .xlsx file"):
        read_excel(path)


@pytest.mark.parametrize("contents", [b"", b"not an Excel workbook"])
def test_invalid_workbook_has_clear_error_and_is_unchanged(tmp_path, contents):
    path = tmp_path / "invalid.xlsx"
    path.write_bytes(contents)
    with pytest.raises(ExcelReadError, match="Could not read .xlsx file") as error:
        read_excel(path)
    assert error.value.__cause__ is not None
    assert path.read_bytes() == contents


def test_zip_without_workbook_is_invalid(tmp_path):
    path = tmp_path / "invalid.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("synthetic.txt", "not a workbook")
    with pytest.raises(ExcelReadError, match="invalid.xlsx"):
        read_excel(path)


def test_permission_error_is_preserved(invoice_file):
    # Deterministic even on systems where tests run with elevated permissions.
    with patch.object(Path, "open", side_effect=PermissionError("Access denied")):
        with pytest.raises(PermissionError, match="Access denied"):
            read_excel(invoice_file)


@pytest.mark.parametrize("headers", [None, ["Vernr", "Utfall"]])
def test_empty_or_header_only_worksheet(tmp_path, headers):
    path = tmp_path / "empty.xlsx"
    workbook = Workbook()
    if headers:
        workbook.active.append(headers)
    workbook.save(path)
    workbook.close()
    result = read_excel(path)
    assert result.empty
    assert result.columns.tolist() == (headers or [])


def test_preserves_duplicate_headers_formulas_and_excel_errors(tmp_path):
    path = tmp_path / "values.XLSX"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Extra", "Extra", None, "Formula", "Error"])
    sheet.append(["001", "NA", "value", "=1+2", "#DIV/0!"])
    workbook.save(path)
    workbook.close()
    result = read_excel(path)
    assert result.columns.tolist() == ["Extra", "Extra", None, "Formula", "Error"]
    assert result.iloc[0].tolist() == ["001", "NA", "value", "=1+2", "#DIV/0!"]
