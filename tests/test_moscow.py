"""Acceptance tests use synthetic invoices only; no reference registers are fabricated."""
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from src.ui_support import analyze_upload
from src.output.report_generator import review_tables

APP = Path(__file__).resolve().parents[1] / 'streamlit_app.py'
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def content(data):
    buffer = BytesIO()
    data.to_excel(buffer, index=False)
    return buffer.getvalue()


def source():
    return pd.DataFrame({'Vertyp': ['X', 'X', 'KR01', 'AAHV90', 'KR01', 'NEW', None],
                         'Konto': [7698, '7699 ', '4000', '4000', ' 7698.0 ', '4000', None],
                         'Vernr': [str(i) for i in range(7)], 'Vrad': [1] * 7,
                         'Verdatum': ['2026-09-03'] * 7, 'Utfall': [10] * 7,
                         'Extra originalkolumn': ['=1+1'] * 7})


def test_partition_original_values_export_and_source_file(tmp_path):
    data = source()
    path = tmp_path / 'original.xlsx'
    path.write_bytes(content(data))
    before, mtime = path.read_bytes(), path.stat().st_mtime_ns
    review = analyze_upload(before)
    kept, excluded = review_tables(review.result.original_data, review.result.filtering)
    assert kept['Vernr'].tolist() == ['5', '6']
    assert excluded['Vernr'].tolist() == ['0', '1', '2', '3', '4']
    restored = pd.concat([kept, excluded.iloc[:, :-1]]).sort_index()
    pd.testing.assert_frame_equal(restored, review.result.original_data)
    assert excluded.iloc[4, -1] == 'Exkluderad – konto 7698; verifikationstyp KR01'
    exported = tmp_path / 'new.xlsx'
    exported.write_bytes(review.downloads['samlad_kontrollfil.xlsx'])
    from openpyxl import load_workbook
    book = load_workbook(exported)
    assert book.sheetnames[:3] == ['Granskning', 'Bortfiltrerade', 'Sammanfattning']
    assert {'Källspårning', 'Källinformation', 'Urvalsmetod', 'Urvalspositioner',
            'Körningsöversikt', 'Exkluderingsregler', 'Registerinformation'} <= set(book.sheetnames)
    assert book['Granskning'].max_row == 3
    assert book['Bortfiltrerade'].max_row == 6
    assert book['Granskning']['G2'].value == '=1+1'
    assert book['Granskning']['G2'].data_type == 's'
    book.close()
    assert path.read_bytes() == before
    assert path.stat().st_mtime_ns == mtime
    assert exported != path
    for detection in review.result.detection_results:
        assert all(check.status == 'NOT_CHECKED' for check in detection.checks)


def test_dynamic_reinclude_additional_type_reset_and_upload_change():
    app = AppTest.from_file(str(APP), default_timeout=20).run()
    app.file_uploader[0].set_value(('invoices.xlsx', content(source()), MIME)).run()
    app.button[0].click().run()
    assert not app.exception
    assert app.button(key="kpi_review").label.split("**")[1] == '2'
    assert 'NEW' in app.multiselect(key='excluded_types').options
    app.multiselect(key='excluded_types').unselect('KR01').run()
    assert not app.exception
    assert app.button(key="kpi_review").label.split("**")[1] == '3'
    assert next(tab for tab in app.tabs if tab.label == 'Granskning').dataframe[0].value['Vernr'].tolist() == ['2', '5', '6']
    review = app.session_state['review']
    assert len(pd.read_excel(BytesIO(review.downloads['granskning.xlsx']))) == 3
    app.multiselect(key='excluded_types').select('NEW').run()
    assert app.button(key="kpi_review").label.split("**")[1] == '2'
    next(b for b in app.button if b.label == 'Återställ filter till standard').click().run()
    assert 'KR01' in app.multiselect(key='excluded_types').value
    assert 'NEW' not in app.multiselect(key='excluded_types').value
    assert app.button(key="kpi_review").label.split("**")[1] == '2'
    messages = [message.value for message in app.info]
    assert 'Upphandlingskontroll – ej tillgänglig. Upphandlingsregister saknas.' in messages
    assert ('Attestkontroll ej genomförd – kräver attestregister eller motsvarande behörighetsunderlag '
            'som inte finns tillgängligt i prototypen.') in messages
    app.file_uploader[0].set_value(('second.xlsx', content(source().iloc[:1]), MIME)).run()
    assert not app.metric
    app.button[0].click().run()
    assert not app.exception
    assert app.button(key="kpi_total").label.split("**")[1] == '1'
    assert app.button(key="kpi_review").label.split("**")[1] == '0'


@pytest.mark.parametrize('missing', ['Konto', 'Vertyp'])
def test_missing_filter_column_has_readable_ui_error(missing):
    review = analyze_upload(content(source().drop(columns=missing)))
    app = AppTest.from_file(str(APP))
    app.session_state['review'] = review
    app.run()
    assert not app.exception
    assert any(missing in error.value and 'saknas' in error.value for error in app.error)
    assert len(review.result.filtering.cleaned_data) + len(review.result.filtering.excluded_data) == 7


def test_header_whitespace_order_and_empty_workbook():
    data = source()
    reordered = data[list(reversed(data.columns))].rename(columns={'Konto': ' Konto ', 'Vertyp': ' Vertyp '})
    first, second = analyze_upload(content(data)), analyze_upload(content(reordered))
    assert first.result.filtering.reasons == second.result.filtering.reasons
    assert list(second.result.original_data.columns) == list(reordered.columns)
    empty = analyze_upload(content(data.iloc[:0]))
    assert not len(empty.result.filtering.cleaned_data)
    assert not len(empty.result.filtering.excluded_data)
    assert empty.downloads['samlad_kontrollfil.xlsx']


def test_metadata_header_upload_filters_rendering_and_export(tmp_path):
    from openpyxl import Workbook
    from src.filtering.filter_engine import load_exclusions
    path = tmp_path / 'metadata_invoices.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.append([None, 'Verifikationslista urval VO', None, None])
    # NEW must remain included even though this source note says otherwise.
    sheet.append([None, 'Tagit bort: VERTYP: AAHV90, AAHV94, NEW', None])
    sheet.append([None] * 12)
    sheet.append(['Kommentar', None, 'Vernr'])
    data = source()
    sheet.append(list(data.columns) + [None, None, None])
    for row in data.itertuples(index=False, name=None):
        sheet.append(list(row) + [None, None, None])
    book.save(path)
    book.close()
    before, mtime = path.read_bytes(), path.stat().st_mtime_ns
    app = AppTest.from_file(str(APP), default_timeout=20).run()
    app.file_uploader[0].set_value(('metadata.xlsx', before, MIME)).run()
    app.button[0].click().run()
    assert not app.exception
    assert not app.error
    assert app.button(key="kpi_total").label.split("**")[1] == '7'
    assert app.button(key="kpi_review").label.split("**")[1] == '2'
    review = app.session_state['review']
    assert {'account', 'verification_type'} <= set(review.result.standardized_data.columns)
    assert review.result.filtering.account_count == 3
    assert review.result.filtering.verification_type_count == 3
    assert app.multiselect(key='excluded_types').value == load_exclusions()['excluded_verification_types']
    assert 'NEW' in app.multiselect(key='excluded_types').options
    assert 'NEW' not in app.multiselect(key='excluded_types').value
    kept = next(tab for tab in app.tabs if tab.label == 'Granskning').dataframe[0].value
    removed = next(tab for tab in app.tabs if tab.label == 'Bortfiltrerade').dataframe[0].value
    assert kept.columns.tolist() == data.columns.tolist()
    assert kept['Vernr'].tolist() == ['5', '6']
    assert removed['Vernr'].tolist() == ['0', '1', '2', '3', '4']
    assert kept.columns.is_unique and removed.columns.is_unique
    with pd.ExcelFile(BytesIO(review.downloads['samlad_kontrollfil.xlsx'])) as exported:
        assert exported.sheet_names[:3] == ['Granskning', 'Bortfiltrerade', 'Sammanfattning']
        assert 'Källspårning' in exported.sheet_names
        assert len(pd.read_excel(exported, sheet_name='Granskning')) == 2
        assert len(pd.read_excel(exported, sheet_name='Bortfiltrerade')) == 5
    assert path.read_bytes() == before
    assert path.stat().st_mtime_ns == mtime


def test_lr_default_filters_ui_population_kpis_and_export():
    from tests.test_feature_package import invoices
    data = invoices(44)
    data['Huvudtext'] = [f'Syntetisk leverantör {i} AB Prelb 1' for i in range(44)]
    data.loc[:3, 'Vertyp'] = ['LR91S', 'LR95S', 'LR91S', 'FBFM']
    data.loc[2, 'Konto'] = '7698'
    original = content(data)
    review = analyze_upload(original, registry_mode='disabled', defer_downloads=True)
    app = AppTest.from_file(str(APP), default_timeout=20)
    app.session_state['review'] = review
    app.run()
    assert not app.exception
    assert {'LR91S', 'LR95S', 'FBFM'} <= set(app.multiselect(key='excluded_types').value)
    assert app.button(key='kpi_total').label.startswith('**44**')
    assert app.button(key='kpi_review').label.startswith('**40**')
    assert app.button(key='kpi_excluded').label.startswith('**4**')
    assert app.button(key='kpi_account').label.startswith('**1**')
    assert app.button(key='kpi_verification_type').label.startswith('**4**')
    assert review.result.summary.eligible_verifications == 40
    assert [v.verification_id for v in review.result.manual_sample] == ['24', '44']
    assert review.result.summary.exclusion_counts['verification_type: LR91S'] == 2
    assert review.result.summary.exclusion_counts['verification_type: LR95S'] == 1
    exported = pd.read_excel(BytesIO(review.downloads['samlad_kontrollfil.xlsx']), sheet_name=None)
    assert len(exported['Granskning']) == 40
    assert len(exported['Bortfiltrerade']) == 4
    assert exported['Bortfiltrerade'].Exkluderingsorsak.iloc[2] == (
        'Exkluderad – konto 7698; verifikationstyp LR91S')
    assert exported['Urvalsmetod']['Antal verifikationer i populationen'].iloc[0] == 40

    for verification_type, count in [('LR91S', 41), ('LR95S', 42)]:
        app.multiselect(key='excluded_types').unselect(verification_type).run()
        assert not app.exception
        changed = app.session_state['review']
        assert changed.result.summary.eligible_verifications == count
        assert app.button(key='kpi_review').label.startswith(f'**{count}**')
        assert len(pd.read_excel(BytesIO(changed.downloads['granskning.xlsx']))) == count
        assert changed.source_content == original
        pd.testing.assert_frame_equal(changed.result.original_data, review.result.original_data)
    assert [v.verification_id for v in changed.result.manual_sample] == ['22', '42']
    assert changed.result.filtering.excluded_data.index.tolist() == [2, 3]
    next(b for b in app.button if b.label == 'Återställ filter till standard').click().run()
    assert not app.exception
    assert {'LR91S', 'LR95S'} <= set(app.multiselect(key='excluded_types').value)
    assert app.session_state['review'].result.summary.eligible_verifications == 40
