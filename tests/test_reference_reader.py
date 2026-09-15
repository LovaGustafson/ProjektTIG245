"""Reference ingestion tests using generic synthetic CSV and Excel files."""

from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import Workbook

from src.ingestion.reference_reader import read_attestation_register, read_supplier_register
from src.models.ingestion import ReadStatus


READERS = [read_supplier_register, read_attestation_register]


@pytest.mark.parametrize("reader", READERS)
@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_generic_load_preserves_source_and_exposes_unconfirmed_schema(tmp_path, reader, suffix):
    path = tmp_path / f"synthetic{suffix}"
    headers = ["arbitrary_code", "counterparty", "unknown_value"]
    records = [["00123", "NA", " synthetic "], ["00456", "NULL", "other"]]
    if suffix == ".xlsx":
        workbook = Workbook()
        workbook.active.append(headers)
        for record in records:
            workbook.active.append(record)
        workbook.save(path)
        workbook.close()
    else:
        path.write_text("arbitrary_code,counterparty,unknown_value\n"
                        "00123,NA, synthetic \n00456,NULL,other\n", encoding="utf-8")
    before, modified = path.read_bytes(), path.stat().st_mtime_ns
    result = reader(path)
    assert result.status == ReadStatus.LOADED
    assert result.register_type == ("supplier_procurement" if reader is read_supplier_register else "attestation")
    assert result.schema_status == "UNCONFIRMED"
    assert result.observed_columns == tuple(headers)
    assert result.todos and all("AK" in todo for todo in result.todos)
    expected = pd.DataFrame(records, columns=pd.Index(headers, dtype=object), dtype=object)
    pd.testing.assert_frame_equal(result.data, expected)
    result.data.iloc[0, 0] = "in-memory edit"
    assert reader(path).data.iloc[0, 0] == "00123"
    assert path.read_bytes() == before
    assert path.stat().st_mtime_ns == modified


@pytest.mark.parametrize("reader", READERS)
def test_missing_reference_file(tmp_path, reader):
    result = reader(tmp_path / "missing.csv")
    assert result.status == ReadStatus.MISSING_FILE
    assert result.data is None
    assert result.schema_status == "UNCONFIRMED"


@pytest.mark.parametrize("reader", READERS)
@pytest.mark.parametrize("suffix,contents", [
    (".xlsx", b"not an Excel file"), (".csv", b""),
    (".csv", b'a,b\n"unterminated,b'), (".csv", b"a,b\n1,2,3\n"),
    (".csv", b"a,b\n1\n"), (".csv", b"a,b\n\xff,2\n"),
])
def test_malformed_reference_files(tmp_path, reader, suffix, contents):
    path = tmp_path / f"malformed{suffix}"
    path.write_bytes(contents)
    result = reader(path)
    assert result.status == ReadStatus.UNREADABLE
    assert result.data is None
    assert result.reason
    assert path.read_bytes() == contents


def test_explicit_csv_delimiter_and_encoding(tmp_path):
    path = tmp_path / "synthetic.csv"
    path.write_bytes("kod;beskrivning\n001;Övrigt\n".encode("cp1252"))
    result = read_attestation_register(path, delimiter=";", encoding="cp1252")
    assert result.status == ReadStatus.LOADED
    assert result.data.iloc[0].tolist() == ["001", "Övrigt"]


def test_csv_duplicate_headers_blank_rows_and_quoted_newlines_are_preserved(tmp_path):
    path = tmp_path / "synthetic.csv"
    path.write_text('\ufeffcode,code,\n001,"two\nlines",NA\n\n002,NULL,\n', encoding="utf-8", newline="")
    result = read_supplier_register(path)
    assert result.status == ReadStatus.LOADED
    assert result.observed_columns == ("code", "code", "")
    assert result.data.values.tolist() == [["001", "two\nlines", "NA"], ["", "", ""], ["002", "NULL", ""]]


def test_explicit_worksheet_selection(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    workbook.active.append(["first"])
    sheet = workbook.create_sheet("Register")
    sheet.append(["generic"])
    sheet.append(["001"])
    workbook.save(path)
    workbook.close()
    result = read_supplier_register(path, sheet_name="Register")
    assert result.status == ReadStatus.LOADED
    assert result.data.iloc[0, 0] == "001"
    assert read_supplier_register(path, sheet_name="missing").status == ReadStatus.UNREADABLE


@pytest.mark.parametrize("reader", READERS)
def test_header_only_csv_loads_without_inventing_required_rows(tmp_path, reader):
    path = tmp_path / "synthetic.csv"
    path.write_text("arbitrary\n", encoding="utf-8")
    result = reader(path)
    assert result.status == ReadStatus.LOADED
    assert result.data.empty
    assert result.schema_status == "UNCONFIRMED"


def test_unsupported_file_type(tmp_path):
    path = tmp_path / "synthetic.json"
    path.write_text("{}")
    assert read_supplier_register(path).status == ReadStatus.UNSUPPORTED


@pytest.mark.parametrize("reference", [None, "", pd.NA])
def test_missing_reference(reference):
    assert read_attestation_register(reference).status == ReadStatus.MISSING_REFERENCE


@pytest.mark.parametrize("reader", READERS)
def test_url_is_never_fetched(reader):
    with patch.object(Path, "open", side_effect=AssertionError("Must not open")):
        assert reader("https://example.invalid/data.csv").status == ReadStatus.UNSUPPORTED


def test_permission_error(tmp_path):
    with patch.object(Path, "open", side_effect=PermissionError("synthetic denial")):
        result = read_supplier_register(tmp_path / "synthetic.csv")
    assert result.status == ReadStatus.UNREADABLE
    assert "denial" in result.reason


def test_directory_is_unreadable(tmp_path):
    path = tmp_path / "directory.csv"
    path.mkdir()
    assert read_attestation_register(path).status == ReadStatus.UNREADABLE
